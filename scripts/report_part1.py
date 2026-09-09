# -*- coding: utf-8 -*-
"""
Report Helpers Module:
- keep_with_next on figure paragraphs prevents separating figures from text.
- Compact cell and paragraph spacing to prevent page overflows and empty pages.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=30, bottom=30, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_academic_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="333333"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="333333"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def set_table_borders(table, color="D0D5DD", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def add_p(doc, text="", style='Normal', space_before=Pt(0), space_after=Pt(4), line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, font_size=Pt(12), color=None, bold_prefix=None, keep_with_next=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = line_spacing
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = font_size
        r_b.bold = True
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = font_size
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
    return p

def add_bullet(doc, text, bold_prefix=None, space_before=Pt(0), space_after=Pt(2.0), line_spacing=1.15):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(12)
        r_b.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_chapter_heading(doc, chap_num, title):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(20)
    p1.paragraph_format.space_after = Pt(6)
    p1.paragraph_format.line_spacing = 1.0
    p1.paragraph_format.keep_with_next = True
    r1 = p1.add_run(f"CHAPTER {chap_num}")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)
    r1.bold = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(14)
    p2.paragraph_format.line_spacing = 1.0
    p2.paragraph_format.keep_with_next = True
    r2 = p2.add_run(title.upper())
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.bold = True

def add_sec_heading(doc, sec_no, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{sec_no} {title.upper()}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    return p

def add_subsec_heading(doc, subsec_no, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{subsec_no} {title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    return p

def add_figure(doc, img_path, fig_no, caption, width=Inches(4.8), height=None, keep_caption_with_next=False):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        if height:
            run_img.add_picture(img_path, height=height)
        else:
            run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = keep_caption_with_next
        r_cap = p_cap.add_run(f"Figure {fig_no}: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.bold = True
    else:
        print(f"Warning: Image not found: {img_path}")

def add_prelim_list(doc, title, headers, rows_data, col_widths, col_alignments=None):
    is_toc = (title == "TABLE OF CONTENTS")
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(12) if is_toc else Pt(16)
    p_hdr.paragraph_format.space_after = Pt(8) if is_toc else Pt(12)
    p_hdr.paragraph_format.keep_with_next = True
    r = p_hdr.add_run(title.upper())
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_margins(hdr_cells[i], top=10 if is_toc else 20, bottom=20 if is_toc else 40, left=30, right=30)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = col_alignments[i] if col_alignments else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0) if is_toc else Pt(1)
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10) if is_toc else Pt(11)
            run.bold = True

    # Content Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        is_chapter = False
        if len(row) >= 3 and row[0].strip().isdigit():
            is_chapter = True
        
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            top_pad = 6 if is_toc else 15
            bot_pad = 6 if is_toc else 15
            set_cell_margins(row_cells[c_idx], top=top_pad, bottom=bot_pad, left=30, right=30)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = col_alignments[c_idx] if col_alignments else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0) if is_toc else Pt(1)
            p.paragraph_format.line_spacing = 1.0 if is_toc else 1.1
            for run in p.runs:
                run.font.name = 'Times New Roman'
                if is_toc:
                    run.font.size = Pt(9.2) if is_chapter else Pt(8.6)
                else:
                    run.font.size = Pt(10.5)
                if is_chapter:
                    run.bold = True

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

def add_table_custom(doc, table_no, caption, headers, rows_data, col_widths=None):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    r_cap = p_cap.add_run(f"Table {table_no}: {caption}")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10.5)
    r_cap.bold = True

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_academic_table_borders(table)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_margins(hdr_cells[i], top=35, bottom=35, left=40, right=40)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            run.bold = True

    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=18, bottom=18, left=40, right=40)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.0)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

def add_code_block(doc, title, code_text):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(f"Listing: {title}")
    r_title.font.name = 'Consolas'
    r_title.font.size = Pt(10)
    r_title.bold = True

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CCCCCC", sz="4")
    cell = table.rows[0].cells[0]
    set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
    set_cell_shading(cell, "F5F6F8")
    cell.width = Inches(6.2)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)

print("report_part1.py compiled with keep_with_next.")

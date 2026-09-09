import io
import re
import os
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from pydantic import BaseModel, EmailStr
from typing import Optional, List

from app.database import get_db
from app.models.user import User, ImportHistory
from app.middleware.auth_middleware import get_current_user
from app.middleware.role_checker import require_role
from app.services.auth_service import hash_password

router = APIRouter(prefix="/admin/faculty-import", tags=["Faculty Import"])

email_regex = re.compile(r"^[^@]+@[^@]+\.[^@]+$")


class FacultyImportItem(BaseModel):
    employee_id: str
    name: str
    email: str
    department: str
    staff_room: Optional[str] = None
    phone_number: Optional[str] = None


class ImportFacultyRequest(BaseModel):
    filename: str
    faculties: List[FacultyImportItem]


@router.get("/template")
async def download_template(
    current_user: dict = Depends(require_role("admin")),
):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Faculty Template"

    headers = [
        "Employee ID",
        "Faculty Name",
        "Email",
        "Department",
        "Staff Room",
        "Phone Number",
    ]
    sheet.append(headers)

    sheet.append(
        ["EMP005", "Dr. Sarah Connor", "sarah@campus.com", "CSE", "Block B, Room 204", "9876543211"]
    )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=faculty_import_template.xlsx"
        },
    )


@router.post("/upload")
async def upload_faculty_excel(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".xlsx", ".xls"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload an Excel sheet (.xlsx or .xls).",
        )

    file_bytes = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read Excel file: {str(e)}",
        )

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Excel sheet is empty or has no faculty rows.",
        )

    # Standardize headers
    first_row = rows[0]
    headers = [
        str(cell).strip().lower().replace(" ", "") if cell else ""
        for cell in first_row
    ]

    header_map = {}
    for idx, h in enumerate(headers):
        if h in ["employeeid", "empid", "id"]:
            header_map["employee_id"] = idx
        elif h in ["facultyname", "name"]:
            header_map["name"] = idx
        elif h in ["email", "emailid"]:
            header_map["email"] = idx
        elif h in ["department", "dept"]:
            header_map["department"] = idx
        elif h in ["staffroom", "room"]:
            header_map["staff_room"] = idx
        elif h in ["phonenumber", "phone", "mobile"]:
            header_map["phone_number"] = idx

    # Check for mandatory headers
    required_keys = ["employee_id", "name", "email", "department"]
    missing_headers = [k for k in required_keys if k not in header_map]
    if missing_headers:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Missing required columns: {', '.join(missing_headers)}",
        )

    valid_records = []
    errors = []
    seen_ids = set()
    seen_emails = set()

    # Query existing DB records to check duplicates
    db_users_res = await db.execute(select(User.employee_id, User.email))
    db_users = db_users_res.all()
    db_ids = {u[0] for u in db_users if u[0]}
    db_emails = {u[1] for u in db_users if u[1]}

    valid_departments = ["CSE", "ECE", "ME", "CE", "IT"]

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        row_data = {}
        for key, col_idx in header_map.items():
            val = row[col_idx]
            row_data[key] = str(val).strip() if val is not None else ""

        emp_id = row_data.get("employee_id", "")
        name = row_data.get("name", "")
        email = row_data.get("email", "")
        dept = row_data.get("department", "").upper()
        staff_room = row_data.get("staff_room", "")
        phone = row_data.get("phone_number", "")

        row_errors = []

        if not emp_id:
            row_errors.append("Employee ID is missing.")
        if not name:
            row_errors.append("Faculty Name is missing.")
        if not email:
            row_errors.append("Email is missing.")
        if not dept:
            row_errors.append("Department is missing.")

        if email and not email_regex.match(email):
            row_errors.append(f"Invalid email format: '{email}'.")
        if dept and dept not in valid_departments:
            row_errors.append(
                f"Invalid department '{dept}'. Must be one of: {', '.join(valid_departments)}."
            )

        if emp_id:
            if emp_id in seen_ids:
                row_errors.append(f"Duplicate Employee ID '{emp_id}' within the Excel sheet.")
            elif emp_id in db_ids:
                row_errors.append(f"Employee ID '{emp_id}' already exists in database.")
            seen_ids.add(emp_id)

        if email:
            if email in seen_emails:
                row_errors.append(f"Duplicate Email '{email}' within the Excel sheet.")
            elif email in db_emails:
                row_errors.append(f"Email '{email}' already registered in database.")
            seen_emails.add(email)

        if row_errors:
            errors.append(
                {
                    "row": row_idx,
                    "employee_id": emp_id,
                    "name": name,
                    "messages": row_errors,
                }
            )
        else:
            valid_records.append(
                {
                    "employee_id": emp_id,
                    "name": name,
                    "email": email,
                    "department": dept,
                    "staff_room": staff_room,
                    "phone_number": phone,
                }
            )

    return {
        "filename": filename,
        "total_records": len(rows) - 1,
        "valid_count": len(valid_records),
        "error_count": len(errors),
        "errors": errors,
        "valid_records": valid_records,
        "isValid": len(errors) == 0,
    }


@router.post("/import")
async def confirm_faculty_import(
    req: ImportFacultyRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    if not req.faculties:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No valid faculty records provided for import.",
        )

    default_pw = "Campus@123"
    hashed_pw = hash_password(default_pw)

    success_count = 0
    fail_count = 0

    for f in req.faculties:
        try:
            # Recheck DB constraints
            dup_res = await db.execute(
                select(User).where((User.employee_id == f.employee_id) | (User.email == f.email))
            )
            if dup_res.scalar_one_or_none():
                fail_count += 1
                continue

            db_user = User(
                name=f.name,
                email=f.email,
                password=hashed_pw,
                role="faculty",
                department=f.department,
                employee_id=f.employee_id,
                staff_room=f.staff_room,
                phone_number=f.phone_number,
                is_first_login=True,
            )
            db.add(db_user)
            success_count += 1
        except Exception:
            fail_count += 1

    # Log import history
    hist = ImportHistory(
        filename=req.filename,
        total_records=len(req.faculties),
        successful_imports=success_count,
        failed_imports=fail_count,
        uploaded_by=current_user["id"],
    )
    db.add(hist)
    await db.commit()

    return {
        "message": f"Successfully imported {success_count} faculty members. {fail_count} failed.",
        "success_count": success_count,
        "fail_count": fail_count,
    }


@router.get("/import-history")
async def get_import_history(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(ImportHistory, User.name)
        .outerjoin(User, ImportHistory.uploaded_by == User.id)
        .order_by(ImportHistory.created_at.desc())
    )
    logs = result.all()

    return [
        {
            "id": log[0].id,
            "filename": log[0].filename,
            "total_records": log[0].total_records,
            "successful_imports": log[0].successful_imports,
            "failed_imports": log[0].failed_imports,
            "uploaded_by_name": log[1] or "System",
            "uploaded_at": log[0].created_at.strftime("%Y-%m-%d %H:%M")
            if log[0].created_at
            else "",
        }
        for log in logs
    ]


@router.get("/export")
async def export_faculty_list(
    department: Optional[str] = None,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    query = select(User).where(User.role == "faculty")
    if department and department.strip() and department.upper() != "ALL":
        query = query.where(User.department == department.strip().upper())
        
    result = await db.execute(query.order_by(User.employee_id))
    faculties = result.scalars().all()

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Faculty List"

    headers = [
        "Employee ID",
        "Faculty Name",
        "Email (Login Username)",
        "Password (Hashed)",
        "Department",
        "Staff Room",
        "Phone Number",
        "Created At",
    ]
    sheet.append(headers)

    for f in faculties:
        sheet.append(
            [
                f.employee_id or "",
                f.name,
                f.email,
                f.password,
                f.department or "",
                f.staff_room or "",
                f.phone_number or "",
                f.created_at.strftime("%Y-%m-%d %H:%M") if f.created_at else "",
            ]
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=faculty_database_export.xlsx"},
    )

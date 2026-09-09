import io
import re
import os
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from pydantic import BaseModel, EmailStr
from typing import Optional, List
from datetime import datetime

from app.database import get_db
from app.models.user import User, ImportHistory
from app.models.communication import Notification
from app.middleware.auth_middleware import get_current_user
from app.middleware.role_checker import require_role
from app.services.auth_service import hash_password

router = APIRouter(prefix="/admin/students", tags=["Student Import"])

email_regex = re.compile(r"^[^@]+@[^@]+\.[^@]+$")


# ─── Pydantic Request Schemas ────────────────────────────────────────────────


class StudentImportItem(BaseModel):
    roll_number: str
    name: str
    email: str
    department: str
    year: str
    section: str
    phone_number: Optional[str] = None


class ImportStudentsRequest(BaseModel):
    filename: str
    students: List[StudentImportItem]


# ─── Helper Functions ────────────────────────────────────────────────────────


def normalize_year(year: str) -> str:
    """Convert numeric year (1,2,3,4) to Roman numeral (I,II,III,IV). Pass-through if already Roman."""
    numeric_to_roman = {"1": "I", "2": "II", "3": "III", "4": "IV"}
    return numeric_to_roman.get(year.strip(), year.strip().upper())


def year_to_semester(year: str) -> int:
    mapping = {"I": 1, "II": 3, "III": 5, "IV": 7}
    return mapping.get(normalize_year(year), 1)


# ─── Endpoints ───────────────────────────────────────────────────────────────


@router.get("/template")
async def download_template(
    current_user: dict = Depends(require_role("admin")),
):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Student Template"

    headers = [
        "Roll Number",
        "Student Name",
        "Email",
        "Department",
        "Year",
        "Section",
        "Phone Number",
    ]
    sheet.append(headers)

    # Add example row
    sheet.append(
        ["22CSR001", "Rahul Sharma", "rahul@campus.com", "CSE", "III", "A", "9876543210"]
    )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=student_import_template.xlsx"
        },
    )


@router.post("/upload")
async def upload_students_excel(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".xlsx", ".xls"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload an Excel sheet (.xlsx or .xls).",
        )

    file_bytes = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read Excel file: {str(e)}",
        )

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Excel sheet is empty or has no student rows.",
        )

    # Standardize headers
    first_row = rows[0]
    headers = [
        str(cell).strip().lower().replace(" ", "") if cell else ""
        for cell in first_row
    ]

    header_map = {}
    for idx, h in enumerate(headers):
        if h in ["rollnumber", "rollno"]:
            header_map["roll_number"] = idx
        elif h in ["studentname", "name"]:
            header_map["name"] = idx
        elif h in ["email", "emailid"]:
            header_map["email"] = idx
        elif h in ["department", "dept"]:
            header_map["department"] = idx
        elif h in ["year"]:
            header_map["year"] = idx
        elif h in ["section"]:
            header_map["section"] = idx
        elif h in ["phonenumber", "phone", "mobile"]:
            header_map["phone_number"] = idx

    # Check for mandatory headers
    required_keys = ["roll_number", "name", "email", "department", "year", "section"]
    missing_headers = [k for k in required_keys if k not in header_map]
    if missing_headers:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Missing required columns: {', '.join(missing_headers)}",
        )

    valid_records = []
    errors = []
    seen_rolls = set()
    seen_emails = set()

    # Query existing DB records to check duplicates
    db_users_res = await db.execute(select(User.roll_number, User.email))
    db_users = db_users_res.all()
    db_rolls = {u[0] for u in db_users if u[0]}
    db_emails = {u[1] for u in db_users if u[1]}

    valid_departments = ["CSE", "ECE", "ME", "CE", "IT"]
    valid_years = ["I", "II", "III", "IV", "1", "2", "3", "4"]

    for row_idx, row in enumerate(rows[1:], start=2):
        # Skip fully empty rows
        if not any(row):
            continue

        row_data = {}
        for key, col_idx in header_map.items():
            val = row[col_idx]
            row_data[key] = str(val).strip() if val is not None else ""

        roll = row_data.get("roll_number", "")
        name = row_data.get("name", "")
        email = row_data.get("email", "")
        dept = row_data.get("department", "").upper()
        year_raw = row_data.get("year", "").strip()
        year = normalize_year(year_raw)
        sect = row_data.get("section", "").upper()
        phone = row_data.get("phone_number", "")

        row_errors = []

        # 1. Missing fields
        if not roll:
            row_errors.append("Roll Number is missing.")
        if not name:
            row_errors.append("Student Name is missing.")
        if not email:
            row_errors.append("Email is missing.")
        if not dept:
            row_errors.append("Department is missing.")
        if not year:
            row_errors.append("Year is missing.")
        if not sect:
            row_errors.append("Section is missing.")

        # 2. Format Validations
        if email and not email_regex.match(email):
            row_errors.append(f"Invalid email format: '{email}'.")
        if dept and dept not in valid_departments:
            row_errors.append(
                f"Invalid department '{dept}'. Must be one of: {', '.join(valid_departments)}."
            )
        if year and year not in ["I", "II", "III", "IV"]:
            row_errors.append(
                f"Invalid year '{year_raw}'. Must be 1-4 or Roman Numeral: I, II, III, or IV."
            )

        # 3. Duplicate checks
        if roll:
            if roll in seen_rolls:
                row_errors.append(f"Duplicate Roll Number '{roll}' within the Excel sheet.")
            elif roll in db_rolls:
                row_errors.append(f"Roll Number '{roll}' already exists in database.")
            seen_rolls.add(roll)

        if email:
            if email in seen_emails:
                row_errors.append(f"Duplicate Email '{email}' within the Excel sheet.")
            elif email in db_emails:
                row_errors.append(f"Email '{email}' already registered in database.")
            seen_emails.add(email)

        if row_errors:
            errors.append(
                {
                    "row": row_idx,
                    "roll_number": roll,
                    "name": name,
                    "messages": row_errors,
                }
            )
        else:
            valid_records.append(
                {
                    "roll_number": roll,
                    "name": name,
                    "email": email,
                    "department": dept,
                    "year": year,
                    "section": sect,
                    "phone_number": phone,
                }
            )

    return {
        "filename": filename,
        "total_records": len(rows) - 1,
        "valid_count": len(valid_records),
        "error_count": len(errors),
        "errors": errors,
        "valid_records": valid_records,
        "isValid": len(errors) == 0,
    }


@router.post("/import")
async def confirm_student_import(
    req: ImportStudentsRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    if not req.students:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No valid student records provided for import.",
        )

    # Configured default password
    default_pw = "Campus@123"
    hashed_pw = hash_password(default_pw)

    success_count = 0
    fail_count = 0

    for s in req.students:
        try:
            # Recheck DB constraints
            dup_res = await db.execute(
                select(User).where((User.roll_number == s.roll_number) | (User.email == s.email))
            )
            if dup_res.scalar_one_or_none():
                fail_count += 1
                continue

            db_user = User(
                name=s.name,
                email=s.email,
                password=hashed_pw,
                role="student",
                department=s.department,
                roll_number=s.roll_number,
                semester=year_to_semester(s.year),
                section=s.section,
                phone_number=s.phone_number,
                is_first_login=True,
                password_changed=False,
            )
            db.add(db_user)
            await db.flush()

            # Create welcome alert / credentials notification for student
            welcome_notif = Notification(
                title="Welcome to CampusOS!",
                message=(
                    f"Welcome {s.name}! Your account has been provisioned. "
                    f"User ID: {s.roll_number} | Default Password: {default_pw}. "
                    "You must reset this password on your first login."
                ),
                category="announcement",
                priority="high",
                user_id=db_user.id,
                created_by=current_user["id"],
            )
            db.add(welcome_notif)
            success_count += 1
        except Exception:
            fail_count += 1

    # Log import history batch
    history_log = ImportHistory(
        filename=req.filename,
        total_records=len(req.students),
        successful_imports=success_count,
        failed_imports=fail_count,
        uploaded_by=current_user["id"],
    )
    db.add(history_log)

    await db.flush()

    return {
        "message": "Bulk student import completed.",
        "total": len(req.students),
        "imported": success_count,
        "failed": fail_count,
    }


@router.get("/import-history")
async def get_import_history(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(ImportHistory, User.name)
        .outerjoin(User, ImportHistory.uploaded_by == User.id)
        .order_by(ImportHistory.created_at.desc())
    )
    logs = result.all()

    return [
        {
            "id": log[0].id,
            "filename": log[0].filename,
            "total_records": log[0].total_records,
            "successful_imports": log[0].successful_imports,
            "failed_imports": log[0].failed_imports,
            "uploaded_by_name": log[1] or "System",
            "uploaded_at": log[0].created_at.strftime("%Y-%m-%d %H:%M")
            if log[0].created_at
            else "",
        }
        for log in logs
    ]


@router.get("/export")
async def export_students_list(
    department: Optional[str] = None,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    query = select(User).where(User.role == "student")
    if department and department.strip() and department.upper() != "ALL":
        query = query.where(User.department == department.strip().upper())
        
    result = await db.execute(query.order_by(User.roll_number))
    students = result.scalars().all()

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Students List"

    headers = [
        "Roll Number",
        "Student Name",
        "Email (Login Username)",
        "Password (Hashed)",
        "Department",
        "Semester",
        "Section",
        "Phone Number",
        "First Login Pending",
        "Created At",
    ]
    sheet.append(headers)

    for s in students:
        sheet.append(
            [
                s.roll_number or "",
                s.name,
                s.email,
                s.password,
                s.department or "",
                s.semester or 1,
                s.section or "A",
                s.phone_number or "",
                "Yes" if s.is_first_login else "No",
                s.created_at.strftime("%Y-%m-%d %H:%M") if s.created_at else "",
            ]
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students_database_export.xlsx"},
    )


@router.get("/stats")
async def get_import_student_stats(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    # Total
    tot_res = await db.execute(
        select(func.count(User.id)).where(User.role == "student")
    )
    total_students = tot_res.scalar() or 0

    # Active (password changed)
    act_res = await db.execute(
        select(func.count(User.id)).where(
            User.role == "student", User.password_changed == True
        )
    )
    active_students = act_res.scalar() or 0

    # Pending first login
    pend_res = await db.execute(
        select(func.count(User.id)).where(
            User.role == "student", User.is_first_login == True
        )
    )
    pending_first = pend_res.scalar() or 0

    # Department wise
    dept_res = await db.execute(
        select(User.department, func.count(User.id))
        .where(User.role == "student")
        .group_by(User.department)
    )
    dept_stats = [{"department": r[0] or "N/A", "count": r[1]} for r in dept_res.all()]

    # Recent imports (last 10 students)
    recent_res = await db.execute(
        select(User)
        .where(User.role == "student")
        .order_by(User.created_at.desc())
        .limit(10)
    )
    recent = recent_res.scalars().all()

    return {
        "total_students": total_students,
        "active_students": active_students,
        "pending_first_login": pending_first,
        "dept_stats": dept_stats,
        "recent_students": [
            {
                "id": s.id,
                "roll_number": s.roll_number,
                "name": s.name,
                "email": s.email,
                "department": s.department,
                "semester": s.semester,
                "section": s.section,
                "created_at": s.created_at.strftime("%Y-%m-%d %H:%M")
                if s.created_at
                else "",
            }
            for s in recent
        ],
    }

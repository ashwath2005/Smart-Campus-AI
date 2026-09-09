# -*- coding: utf-8 -*-
"""
Chapter 4 Builder Module for Smart Campus AI (SCME-AWN)
Contains architectural topology, DFDs, Database Schema, and all 10 Algorithm Formulations.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from report_part1 import (
    set_cell_margins, set_cell_shading, set_table_borders,
    add_p, add_bullet, add_chapter_heading, add_sec_heading,
    add_subsec_heading, add_figure, add_table_custom
)
from data_algorithms import ALGORITHMS_DATA

def build_chapter_4(doc):
    doc.add_page_break()
    add_chapter_heading(doc, 4, "SYSTEM ARCHITECTURE AND DESIGN")

    add_sec_heading(doc, "4.1", "High-Level Architectural Topology (GACB Framework)")
    add_p(doc,
        "The Smart Campus Management Ecosystem (SCME-AWN) is engineered around a modern, decoupled asynchronous microservices "
        "topology. The architecture comprises four interconnected computational tiers: (1) Client Presentation Tier, (2) API Gateway "
        "and Security Layer, (3) Asynchronous Application & Algorithmic Engine, and (4) Relational Persistence & Cloud AI Services."
    )
    add_p(doc,
        "The flagship architectural innovation of SCME-AWN is the Generative AI Context Binding (GACB) framework. Traditional AI "
        "chatbots either guess from static pre-training weights (causing severe hallucinations) or query unindexed vector embeddings "
        "that become stale when operational ERP states change. GACB solves this problem by introducing a dynamic, zero-indexing "
        "Context Compiler. When a user submits an administrative or academic query, the Context Compiler executes parameterized "
        "relational queries against live PostgreSQL/SQLite tables, serializes current state tuples (active lecture slots, attendance "
        "records, approved leave slips, official notices) into compact JSON structures, and binds them into the LLM system prompt. "
        "The LLM reasons directly over authoritative institutional truth, delivering zero-hallucination responses in real-time."
    )

    # Embed Block Diagram
    add_figure(doc, os.path.join('Report', 'images', 'block_diagram.png'), "4.1", "Block Diagram of Smart Campus AI Management Ecosystem", width=Inches(4.5))
    add_p(doc, "Figure 4.1 illustrates the architectural block diagram of SCME-AWN, detailing the multi-tiered decoupling between client presentation interfaces, FastAPI ASGI gateway, context compiler, and relational persistence storage.", space_after=Pt(3), line_spacing=1.14)

    # Embed Flow Diagram
    add_figure(doc, os.path.join('Report', 'images', 'flow_diagram.png'), "4.2", "High-Level System Flow Diagram of Campus Operations", width=Inches(4.5))
    add_p(doc, "Figure 4.2 charts the sequential operational flow across authentication, dynamic context compilation, Gemini LLM reasoning, and multi-channel notification dispatching.", space_after=Pt(3), line_spacing=1.14)

    add_sec_heading(doc, "4.2", "Data Flow Diagrams (Level 0, Level 1, Level 2)")
    add_p(doc,
        "The flow of data across system boundaries is modeled hierarchically through three levels of Data Flow Diagrams (DFDs):"
    )
    add_subsec_heading(doc, "4.2.1", "DFD Level 0: Context Diagram")
    add_p(doc,
        "The Level 0 Context Diagram models SCME-AWN as a centralized process interacting with four external entities: (1) Students, "
        "(2) Faculty Members, (3) Campus Administrators, and (4) External Cloud AI & Notification APIs (Google Gemini, Firebase FCM). "
        "Students submit queries, view personalized dashboards, submit gate passes, and log telemetry interactions. Faculty log attendance, "
        "approve leaves, and manage course schedules. Administrators control user permissions, trigger timetable optimization, and "
        "reallocate physical facilities. The system outputs real-time timetable alerts, personalized learning recommendations, and verified "
        "digital gate passes."
    )

    add_subsec_heading(doc, "4.2.2", "DFD Level 1: Functional Subsystem Decomposition")
    add_p(doc,
        "The Level 1 DFD decomposes the system into six core operational subsystems: (1) Authentication & RBAC Engine, (2) Timetable "
        "& Facility Scheduler, (3) Attendance & Gate Pass Manager, (4) AI Learning Intelligence Engine (CLPA/KDPA), (5) GACB Prompt "
        "Compiler & LLM Bridge, and (6) Wearable Push Notification Dispatcher. Data flows bidirectionally between these functional "
        "processes and the centralized relational database stores."
    )

    add_subsec_heading(doc, "4.2.3", "DFD Level 2: GACB Engine & Telemetry Pipeline")
    add_p(doc,
        "The Level 2 DFD provides a granular inspection of the GACB pipeline and telemetry tracking loops. User queries pass through "
        "a JWT Security Validator to extract user ID and role permissions. The SQL Context Compiler retrieves relevant domain records "
        "(e.g., student's timetable for today, enrolled course attendance, approved leave logs). The System Prompt Assembler formats "
        "the context JSON, attaches safety instructions, and submits the payload to Google Gemini 2.5 Flash via asynchronous HTTP/2 "
        "connections. Concurrently, user interaction telemetry (scroll depth, quiz scores, code execution) is captured by the CLPA "
        "pipeline, updating the cognitive profile vector asynchronously in background worker queues."
    )

    add_sec_heading(doc, "4.3", "Relational Database Schema & Entity Summary")
    add_p(doc,
        "SCME-AWN's persistent state is architected around a highly normalized relational schema implemented in PostgreSQL (with SQLite "
        "fallback). The schema enforces strict foreign key constraints, cascading deletions, and indexed query optimizations for high-concurrency "
        "lookups. The core database entities are summarized in Table 4.1."
    )

    db_headers = ["Entity / Table", "Primary Key", "Key Attributes & Foreign Keys", "System Function & Description"]
    db_rows = [
        ["users", "id (UUID)", "email, password_hash, full_name, role, reg_no, dept_id", "Central identity store with role-based attributes and Bcrypt hashes"],
        ["departments", "id (Int)", "code, name, building, floor_count", "Academic department divisions and physical campus building mappings"],
        ["classrooms", "id (Int)", "room_number, building, capacity, has_projector, has_ac, has_lab", "Physical classroom inventory with equipment flags for DCRA+ allocation"],
        ["courses", "id (Int)", "code, title, credits, semester, dept_id, syllabus_dag_json", "Course catalog with credit hours and prerequisite syllabus graph structures"],
        ["timetable_slots", "id (Int)", "day_of_week, slot_index, course_id, faculty_id, room_id, section", "Master schedule slots linking courses, instructors, sections, and rooms"],
        ["attendance_records", "id (Int)", "student_id, course_id, date, slot_index, status (P/A/OD)", "Daily lecture attendance logs tracking physical presence and duty leaves"],
        ["gate_passes", "id (Int)", "student_id, reason, departure_time, expected_return, status, qr_token", "Digital gate pass records with parent OTP and security checkpoint states"],
        ["student_learning_interactions", "id (Int)", "student_id, archetype_scores_json, dwell_time, streak_days", "CLPA cognitive telemetry tracking learning styles and daily active streaks"],
        ["topic_retention_records", "id (Int)", "student_id, topic_id, last_reviewed_at, review_count, stability_s", "KDPA memory retention tracking applying modified Ebbinghaus decay math"],
        ["roadmaps", "id (Int)", "student_id, subject_id, target_exam_date, priority_score, tasks_json", "ALRA dynamic study milestones balancing exam countdowns and attendance gaps"],
        ["question_evolutions", "id (Int)", "topic_id, question_text, options_json, correct_idx, bloom_tier, qes", "ICQEA diagnostic questions with Bloom's taxonomy and difficulty ratings"],
        ["notifications", "id (Int)", "user_id, title, message, priority (1-4), is_read, dispatched_at", "Push notification queue with priority levels for wearable and mobile clients"]
    ]
    add_table_custom(doc, "4.1", "Relational Database Schema & Entity Summary", db_headers, db_rows)

    add_sec_heading(doc, "4.4", "System Algorithms & Mathematical Formulation")
    add_p(doc,
        "SCME-AWN incorporates ten proprietary algorithmic formulations to automate academic planning, resource optimization, "
        "and proactive student interventions. Each algorithm is detailed below with its mathematical formulation, computational complexity, "
        "and execution workflow."
    )

    for algo in ALGORITHMS_DATA:
        add_subsec_heading(doc, algo['no'], algo['name'])
        add_p(doc, f"Operational Module: {algo['module']}", bold=True, italic=True)
        add_p(doc, algo['objective'], bold_prefix="Objective: ")
        add_p(doc, algo['math_desc'], bold_prefix="Mathematical Formulation:\n")
        add_p(doc, algo['complexity'], bold_prefix="Computational Complexity: ")
        add_p(doc, "Step-by-Step Execution Workflow:", bold=True)
        for step in algo['workflow']:
            add_bullet(doc, step)

    # CLPA Weights Table
    clpa_headers = ["Telemetry Metric", "Factor Symbol", "Weight Coefficient", "Cognitive Interpretation"]
    clpa_rows = [
        ["Code Execution Attempts & Success Rate", "E_code", "0.35", "Indicates high Practical / Hands-on learning affinity"],
        ["Textbook / Document Dwell Time & Scroll Depth", "E_read", "0.25", "Indicates Reading / Comprehensive theoretical absorption style"],
        ["Diagnostic Assessment & Practice Quiz Accuracy", "E_quiz", "0.25", "Reflects Analytical / Critical thinking mastery"],
        ["Daily Platform Streak & Regularity Index", "E_streak", "0.15", "Demonstrates Consistent / Disciplined study habits"]
    ]
    add_table_custom(doc, "4.2", "CLPA Learning Style Scoring Coefficients & Telemetry Factors", clpa_headers, clpa_rows)

    # DCRA+ Weights Table
    dcra_headers = ["Evaluation Metric", "Factor Code", "Importance Weight (w_i)", "Operational Significance"]
    dcra_rows = [
        ["Capacity Fit", "f_1", "0.25", "Ensures classroom capacity matches or slightly exceeds section size without oversubscription"],
        ["Projector / AV Readiness", "f_2", "0.15", "Mandatory requirement for modern multi-media and slide-based lecture delivery"],
        ["Proximity to Original Room", "f_3", "0.15", "Minimizes student and faculty transit walking distance during sudden room shifts"],
        ["Air Conditioning / Ventilation", "f_4", "0.10", "Maintains thermal comfort standards for large student cohorts"],
        ["Elevator / Accessibility Proximity", "f_5", "0.10", "Accommodates students and faculty with physical accessibility requirements"],
        ["Lab Workstation Hardware", "f_6", "0.08", "Evaluates availability of dedicated computers and test equipment for practical sessions"],
        ["Acoustic Isolation Index", "f_7", "0.05", "Selects rooms shielded from external campus traffic and construction noise"],
        ["Department Wing Bonus", "f_8", "0.04", "Prioritizes rooms located within the host department's academic block"],
        ["Laptop Power Outlet Density", "f_9", "0.04", "Favors classrooms with peripheral charging points for BYOD engineering sessions"],
        ["Interactive Smart Board", "f_10", "0.04", "Favors rooms equipped with stylus digital touch displays for mathematical derivations"]
    ]
    add_table_custom(doc, "4.3", "DCRA+ 10-Factor Suitability Weights Matrix", dcra_headers, dcra_rows)

    add_sec_heading(doc, "4.5", "Security Architecture & Role-Based Access Control")
    add_p(doc,
        "Given that educational ERPs house highly sensitive academic, parental, and disciplinary records, SCME-AWN implements "
        "a zero-trust security framework. User authentication employs stateless JSON Web Tokens (JWT) signed with HMAC-SHA256 "
        "cryptographic secrets and configured with short expiry windows (15 minutes for access tokens, 7 days for refresh tokens). "
        "Passwords are salt-hashed using the Bcrypt algorithm with a work factor of 12 rounds. Role-Based Access Control (RBAC) "
        "is strictly enforced at both API gateway middleware and database repository layers, ensuring students can only access "
        "their own personal records, faculty can only grade their assigned course sections, and administrative functions remain "
        "strictly restricted to authorized credentials. All database interactions utilize parameterized SQLAlchemy ORM statements, "
        "eliminating SQL injection vulnerabilities entirely.",
        space_after=Pt(3), line_spacing=1.14
    )

    add_subsec_heading(doc, "4.5.1", "Cryptographic Authentication & Token Lifecycle Management")
    add_p(doc,
        "User identification is established via an OAuth2 Password Bearer flow with dual-token issuance. Upon successful authentication, "
        "the FastAPI security handler generates a short-lived access token embedded with role claims, user ID, and issue timestamp, alongside "
        "a cryptographically random refresh token. To prevent session hijacking across distributed client devices, refresh tokens are stored "
        "in an in-memory Redis token revocation cache with absolute expiry limits. Revoked or logged-out tokens are immediately blacklisted, "
        "ensuring instant session invalidation without waiting for access token expiration.",
        space_after=Pt(3), line_spacing=1.14
    )

    add_subsec_heading(doc, "4.5.2", "Role-Based Access Control (RBAC) Policy Matrix")
    add_p(doc,
        "Authorization policies are enforced using FastAPI's dependency injection system (`SecurityScopes`). Four primary operational tiers "
        "govern access rights: (1) Student Tier: Read access to personal attendance, enrolled timetables, syllabus roadmaps, and gate pass "
        "submission; (2) Faculty Tier: Scoped write access to assigned section attendance, grade submissions, leave applications, and study material uploads; "
        "(3) Warden Tier: Approval authority over digital hostel gate passes and parent OTP verification telemetry; (4) Administrator Tier: Unrestricted "
        "system auditing, user provisioning, global timetable regeneration, and classroom facility overrides.",
        space_after=Pt(3), line_spacing=1.14
    )

    add_subsec_heading(doc, "4.5.3", "Privacy Preservation & AI Context Sanitization (FERPA / GDPR)")
    add_p(doc,
        "A critical vulnerability in enterprise generative AI integration is the accidental exposure of Personally Identifiable Information (PII) "
        "to external model APIs. The GACB Context Compiler implements automated sanitization layers before compiling system prompts. "
        "Student phone numbers, parent contact details, hashed authentication credentials, and disciplinary remarks are strictly stripped "
        "during relational state extraction. Only aggregated operational metadata necessary to answer the contextual query is bound into "
        "the prompt, guaranteeing complete alignment with institutional data privacy standards.",
        space_after=Pt(3), line_spacing=1.14
    )

print("Chapter 4 compiled.")

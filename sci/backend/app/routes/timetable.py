import os
import json
import io
import csv
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, update, delete, func
from pydantic import BaseModel
from typing import Optional, List
from datetime import time, datetime

from app.database import get_db
from app.models.user import Timetable, TimetableEntry, TimetableVersion, User, ImportHistory
from app.models.communication import Notification
from app.models.academic import Classroom, ClassroomAllocation, ReallocationLog, Section
from app.middleware.auth_middleware import get_current_user
from app.middleware.role_checker import require_role
from app.utils.file_upload import save_upload_file, BACKEND_DIR
from app.services.gemini_service import parse_timetable_from_text, parse_timetable_from_image
from app.algorithms.dcra_service import (
    DCRA_Algorithm,
    DynamicReallocationEngine,
    ClassroomHealthIndexCalculator,
    SuitabilityIndexCalculator,
    DEFAULT_WEIGHTS
)

router = APIRouter(prefix="/timetable", tags=["Timetable"])


# ─── Helper Functions ────────────────────────────────────────────────────────


def map_semester_to_year(semester: int) -> str:
    if semester in (1, 2):
        return "I"
    elif semester in (3, 4):
        return "II"
    elif semester in (5, 6):
        return "III"
    elif semester in (7, 8):
        return "IV"
    return "I"


def time_to_minutes(t) -> int:
    if isinstance(t, str):
        parts = t.split(":")
        return int(parts[0]) * 60 + int(parts[1])
    elif isinstance(t, time):
        return t.hour * 60 + t.minute
    return 0


# ─── Pydantic Request Models ──────────────────────────────────────────────────


class TimetableEntrySchema(BaseModel):
    day: str
    subject: str
    faculty: Optional[str] = None
    start_time: str  # HH:MM
    end_time: str  # HH:MM
    room: Optional[str] = None


class ConfirmTimetableRequest(BaseModel):
    department: str
    year: str
    section: str
    file_path: str
    entries: List[TimetableEntrySchema]


# ─── Endpoints ───────────────────────────────────────────────────────────────


@router.get("/my")
async def get_my_timetable(
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if current_user.get("role") == "faculty":
        entries_res = await db.execute(
            select(TimetableEntry)
            .join(Timetable, TimetableEntry.timetable_id == Timetable.id)
            .where(
                TimetableEntry.faculty == current_user.get("name"),
                Timetable.is_active == True
            )
            .order_by(TimetableEntry.day, TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()
    else:
        year = map_semester_to_year(current_user.get("semester", 1))
        dept = current_user.get("department")
        section = current_user.get("section", "A")

        result = await db.execute(
            select(Timetable)
            .where(
                Timetable.department == dept,
                Timetable.year == year,
                Timetable.section == section,
                Timetable.is_active == True,
            )
        )
        timetable = result.scalar_one_or_none()

        if not timetable:
            return []

        entries_res = await db.execute(
            select(TimetableEntry)
            .where(TimetableEntry.timetable_id == timetable.id)
            .order_by(TimetableEntry.day, TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()

    return [
        {
            "id": e.id,
            "day": e.day,
            "startTime": e.start_time.strftime("%H:%M") if e.start_time else "",
            "endTime": e.end_time.strftime("%H:%M") if e.end_time else "",
            "subject": e.subject,
            "subjectName": e.subject,
            "faculty": e.faculty or "N/A",
            "room": e.room or "N/A",
        }
        for e in entries
    ]


@router.get("/today")
async def get_today_classes(
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    today_name = datetime.now().strftime("%A")
    if current_user.get("role") == "faculty":
        entries_res = await db.execute(
            select(TimetableEntry)
            .join(Timetable, TimetableEntry.timetable_id == Timetable.id)
            .where(
                TimetableEntry.faculty == current_user.get("name"),
                TimetableEntry.day == today_name,
                Timetable.is_active == True
            )
            .order_by(TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()
    else:
        year = map_semester_to_year(current_user.get("semester", 1))
        dept = current_user.get("department")
        section = current_user.get("section", "A")

        result = await db.execute(
            select(Timetable)
            .where(
                Timetable.department == dept,
                Timetable.year == year,
                Timetable.section == section,
                Timetable.is_active == True,
            )
        )
        timetable = result.scalar_one_or_none()

        if not timetable:
            return []

        entries_res = await db.execute(
            select(TimetableEntry)
            .where(
                TimetableEntry.timetable_id == timetable.id,
                TimetableEntry.day == today_name,
            )
            .order_by(TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()

    return [
        {
            "id": e.id,
            "day": e.day,
            "startTime": e.start_time.strftime("%H:%M") if e.start_time else "",
            "endTime": e.end_time.strftime("%H:%M") if e.end_time else "",
            "subject": e.subject,
            "subjectName": e.subject,
            "faculty": e.faculty or "N/A",
            "room": e.room or "N/A",
        }
        for e in entries
    ]


@router.get("/next")
async def get_next_class(
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    today_name = datetime.now().strftime("%A")
    now_time = datetime.now().time()

    if current_user.get("role") == "faculty":
        entries_res = await db.execute(
            select(TimetableEntry)
            .join(Timetable, TimetableEntry.timetable_id == Timetable.id)
            .where(
                TimetableEntry.faculty == current_user.get("name"),
                TimetableEntry.day == today_name,
                Timetable.is_active == True
            )
            .order_by(TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()
    else:
        year = map_semester_to_year(current_user.get("semester", 1))
        dept = current_user.get("department")
        section = current_user.get("section", "A")

        result = await db.execute(
            select(Timetable)
            .where(
                Timetable.department == dept,
                Timetable.year == year,
                Timetable.section == section,
                Timetable.is_active == True,
            )
        )
        timetable = result.scalar_one_or_none()

        if not timetable:
            return None

        entries_res = await db.execute(
            select(TimetableEntry)
            .where(
                TimetableEntry.timetable_id == timetable.id,
                TimetableEntry.day == today_name,
            )
            .order_by(TimetableEntry.start_time)
        )
        entries = entries_res.scalars().all()

    next_entry = None
    for e in entries:
        if e.start_time > now_time:
            next_entry = e
            break

    if not next_entry:
        return None

    return {
        "id": next_entry.id,
        "day": next_entry.day,
        "startTime": next_entry.start_time.strftime("%H:%M") if next_entry.start_time else "",
        "endTime": next_entry.end_time.strftime("%H:%M") if next_entry.end_time else "",
        "subject": next_entry.subject,
        "subjectName": next_entry.subject,
        "faculty": next_entry.faculty or "N/A",
        "room": next_entry.room or "N/A",
    }


@router.post("/upload")
async def upload_timetable(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".pdf", ".png", ".jpg", ".jpeg", ".xls", ".xlsx"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload PDF, Excel, or image (PNG, JPG, JPEG).",
        )

    # Save the file to unique path
    relative_path = await save_upload_file(file, "timetables")
    absolute_path = os.path.join(BACKEND_DIR, relative_path)

    # Read bytes from file
    with open(absolute_path, "rb") as f:
        file_bytes = f.read()

    parsed_data = None
    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(absolute_path)
            text = ""
            for page in doc:
                text += page.get_text()
            doc.close()
            parsed_data = await parse_timetable_from_text(text)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Error extracting PDF text: {str(e)}",
            )
    elif ext in [".xls", ".xlsx"]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(absolute_path, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    if any(val is not None for val in row):
                        row_str = " | ".join(str(val) if val is not None else "" for val in row)
                        lines.append(row_str)
            text = "\n".join(lines)
            parsed_data = await parse_timetable_from_text(text)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Error parsing Excel timetable: {str(e)}",
            )
    else:
        mime_type = file.content_type or "image/png"
        try:
            parsed_data = await parse_timetable_from_image(file_bytes, mime_type)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Error parsing timetable image: {str(e)}",
            )

    if not parsed_data:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to parse timetable data from Gemini.",
        )

    parsed_dept = parsed_data.get("department", "CSE")
    parsed_year = parsed_data.get("year", "II")
    parsed_section = parsed_data.get("section", "A")
    parsed_entries = parsed_data.get("entries", [])

    # Run validations
    conflicts = []
    entries_with_minutes = []

    for entry in parsed_entries:
        day = entry.get("day")
        st = entry.get("start_time")
        et = entry.get("end_time")
        if not day or not st or not et:
            continue
        try:
            s_min = time_to_minutes(st)
            e_min = time_to_minutes(et)
            entries_with_minutes.append((day, s_min, e_min, entry))
        except Exception:
            continue

    # 1. Internal self-collisions
    for i in range(len(entries_with_minutes)):
        for j in range(i + 1, len(entries_with_minutes)):
            day1, s1, e1, entry1 = entries_with_minutes[i]
            day2, s2, e2, entry2 = entries_with_minutes[j]
            if day1.lower() == day2.lower() and max(s1, s2) < min(e1, e2):
                conflicts.append(
                    {
                        "type": "internal",
                        "message": f"Internal overlap: '{entry1.get('subject')}' and '{entry2.get('subject')}' on {day1} ({entry1.get('start_time')}-{entry1.get('end_time')} vs {entry2.get('start_time')}-{entry2.get('end_time')})",
                        "entry1": entry1,
                        "entry2": entry2,
                    }
                )

    # 2. Database active collisions (Room, Faculty)
    db_res = await db.execute(
        select(TimetableEntry, Timetable)
        .join(Timetable)
        .where(Timetable.is_active == True)
    )
    all_db_entries = db_res.all()

    # Filter out active timetable for the same target class (department, year, section)
    other_db_entries = [
        (e, t)
        for e, t in all_db_entries
        if not (
            t.department == parsed_dept
            and t.year == parsed_year
            and t.section == parsed_section
        )
    ]

    for day, s_min, e_min, entry in entries_with_minutes:
        room = entry.get("room")
        faculty = entry.get("faculty")

        for db_entry, db_timetable in other_db_entries:
            if db_entry.day.lower() != day.lower():
                continue

            db_s = time_to_minutes(db_entry.start_time)
            db_e = time_to_minutes(db_entry.end_time)

            if max(s_min, db_s) < min(e_min, db_e):
                if (
                    room
                    and db_entry.room
                    and room.strip().lower() == db_entry.room.strip().lower()
                ):
                    conflicts.append(
                        {
                            "type": "room",
                            "message": f"Room overlap: '{room}' is booked by {db_timetable.department} {db_timetable.year}-{db_timetable.section} for '{db_entry.subject}' on {day} ({db_entry.start_time.strftime('%H:%M')}-{db_entry.end_time.strftime('%H:%M')})",
                            "entry1": entry,
                            "entry2": {
                                "subject": db_entry.subject,
                                "faculty": db_entry.faculty,
                                "room": db_entry.room,
                                "day": db_entry.day,
                                "start_time": db_entry.start_time.strftime("%H:%M"),
                                "end_time": db_entry.end_time.strftime("%H:%M"),
                                "timetable": f"{db_timetable.department} {db_timetable.year}-{db_timetable.section}",
                            },
                        }
                    )
                if (
                    faculty
                    and db_entry.faculty
                    and faculty.strip().lower() == db_entry.faculty.strip().lower()
                ):
                    conflicts.append(
                        {
                            "type": "faculty",
                            "message": f"Faculty conflict: '{faculty}' is already teaching {db_timetable.department} {db_timetable.year}-{db_timetable.section} '{db_entry.subject}' on {day} ({db_entry.start_time.strftime('%H:%M')}-{db_entry.end_time.strftime('%H:%M')})",
                            "entry1": entry,
                            "entry2": {
                                "subject": db_entry.subject,
                                "faculty": db_entry.faculty,
                                "room": db_entry.room,
                                "day": db_entry.day,
                                "start_time": db_entry.start_time.strftime("%H:%M"),
                                "end_time": db_entry.end_time.strftime("%H:%M"),
                                "timetable": f"{db_timetable.department} {db_timetable.year}-{db_timetable.section}",
                            },
                        }
                    )

    return {
        "file_path": relative_path,
        "department": parsed_dept,
        "year": parsed_year,
        "section": parsed_section,
        "entries": parsed_entries,
        "isValid": len(conflicts) == 0,
        "conflicts": conflicts,
    }


from app.controllers.timetable_controller import TimetableController
from app.dtos.timetable_dto import GenerateTimetableRequest as AutoGenerateTimetableRequest


@router.post("/generate")
async def generate_timetable(
    req: AutoGenerateTimetableRequest,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await TimetableController.generate_timetable(req, db, current_user)




@router.post("/publish")
async def publish_timetable(
    req: dict,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    timetable_id = req.get("timetable_id")
    if not timetable_id:
        raise HTTPException(status_code=400, detail="timetable_id is required")
    return await TimetableController.publish_timetable(timetable_id, db, current_user)


@router.delete("/regenerate")
async def regenerate_timetable(
    timetable_id: int,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await TimetableController.regenerate_timetable(timetable_id, db, current_user)



@router.post("/confirm")
async def confirm_timetable(
    req: ConfirmTimetableRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    # 1. Deactivate old active timetables
    await db.execute(
        update(Timetable)
        .where(
            Timetable.department == req.department,
            Timetable.year == req.year,
            Timetable.section == req.section,
            Timetable.is_active == True,
        )
        .values(is_active=False)
    )

    # 2. Create new active timetable
    new_tt = Timetable(
        department=req.department,
        year=req.year,
        section=req.section,
        is_active=True,
    )
    db.add(new_tt)
    await db.flush()

    # 3. Create entries
    for entry in req.entries:
        try:
            start = datetime.strptime(entry.start_time, "%H:%M").time()
            end = datetime.strptime(entry.end_time, "%H:%M").time()
        except ValueError:
            try:
                start = datetime.strptime(entry.start_time, "%H:%M:%S").time()
                end = datetime.strptime(entry.end_time, "%H:%M:%S").time()
            except ValueError:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid time format for {entry.subject}. Use HH:MM.",
                )

        db_entry = TimetableEntry(
            timetable_id=new_tt.id,
            subject=entry.subject,
            faculty=entry.faculty,
            room=entry.room,
            day=entry.day,
            start_time=start,
            end_time=end,
        )
        db.add(db_entry)
        await db.flush() # get id

        is_lunch = entry.subject.lower() in ["lunch", "lunch break", "break", "interval"]
        if entry.room and entry.room != "TBD" and entry.room != "N/A" and not is_lunch:
            c_stmt = select(Classroom).where(Classroom.room_number == entry.room)
            c_res = await db.execute(c_stmt)
            classroom_obj = c_res.scalar_one_or_none()
            if classroom_obj:
                chi = ClassroomHealthIndexCalculator.calculate_chi(classroom_obj)
                
                # Fetch class size
                year_sems = {"I": [1, 2], "II": [3, 4], "III": [5, 6], "IV": [7, 8]}
                sems = year_sems.get(req.year, [1, 2])
                stud_stmt = select(func.count(User.id)).where(
                    User.role == "student",
                    User.department == req.department,
                    User.section == req.section,
                    User.semester.in_(sems)
                )
                stud_res = await db.execute(stud_stmt)
                class_size = stud_res.scalar() or 45

                needs_lab = any(kw in entry.subject.lower() for kw in ["lab", "practical", "workshop", "seminar"])
                suitability = SuitabilityIndexCalculator.calculate_suitability(
                    classroom_obj, chi, class_size, class_size, needs_lab,
                    entry.subject, req.department, 0.0, 0.0, 0.0
                )
                
                ca = ClassroomAllocation(
                    classroom_id=classroom_obj.id,
                    timetable_entry_id=db_entry.id,
                    suitability_score=suitability["score"],
                    explanation=json.dumps({
                        "capacity_match": f"100% ({class_size} students in room capacity {classroom_obj.capacity})",
                        "equipment_compatibility": "100% Match",
                        "movement_score": "Excellent",
                        "predicted_occupancy": f"{class_size} students",
                        "future_conflict_risk": "Low",
                        "classroom_health": f"{chi['score']}% ({chi['status']})",
                        "overall_suitability": f"{suitability['score']}% - Confirmed"
                    }),
                    predicted_occupancy=class_size,
                    occupancy_confidence=1.0,
                    status="active"
                )
                db.add(ca)

    # 4. Save version record
    version_res = await db.execute(
        select(TimetableVersion)
        .join(Timetable)
        .where(
            Timetable.department == req.department,
            Timetable.year == req.year,
            Timetable.section == req.section,
        )
    )
    versions = version_res.scalars().all()
    next_version = len(versions) + 1

    parsed_json_str = json.dumps(
        {
            "department": req.department,
            "year": req.year,
            "section": req.section,
            "entries": [e.dict() for e in req.entries],
        }
    )

    new_version = TimetableVersion(
        timetable_id=new_tt.id,
        version=next_version,
        file_url=req.file_path,
        parsed_data=parsed_json_str,
        uploaded_by=current_user["id"],
    )
    db.add(new_version)

    # 5. Send Notification
    year_map = {"I": 1, "II": 2, "III": 3, "IV": 4}
    target_yr = year_map.get(req.year, 1)

    notification = Notification(
        title=f"New Timetable Released — {req.department} Year {req.year}-{req.section}",
        message=f"The active timetable for {req.department} Year {req.year} Section {req.section} has been updated. Please check your schedule.",
        category="academic",
        priority="normal",
        target_role="student",
        department=req.department,
        target_year=target_yr,
        created_by=current_user["id"],
    )
    db.add(notification)

    await db.flush()
    return {
        "message": "Timetable confirmed and distributed successfully",
        "timetable_id": new_tt.id,
    }


@router.get("/list")
async def list_timetables(
    department: Optional[str] = None,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    query = select(Timetable)
    if department and department.strip() and department.upper() != "ALL":
        query = query.where(Timetable.department == department.strip().upper())

    result = await db.execute(
        query.order_by(
            Timetable.department,
            Timetable.year,
            Timetable.section,
            Timetable.created_at.desc(),
        )
    )
    timetables = result.scalars().all()

    response = []
    for tt in timetables:
        v_res = await db.execute(
            select(TimetableVersion, User)
            .outerjoin(User, TimetableVersion.uploaded_by == User.id)
            .where(TimetableVersion.timetable_id == tt.id)
        )
        versions_data = v_res.all()

        versions_list = []
        for v, user in versions_data:
            versions_list.append(
                {
                    "id": v.id,
                    "version": v.version,
                    "file_url": v.file_url,
                    "uploaded_by_name": user.name if user else "System",
                    "uploaded_at": v.uploaded_at.strftime("%Y-%m-%d %H:%M")
                    if v.uploaded_at
                    else "",
                }
            )

        entries_res = await db.execute(
            select(TimetableEntry).where(TimetableEntry.timetable_id == tt.id)
        )
        entries = entries_res.scalars().all()

        response.append(
            {
                "id": tt.id,
                "department": tt.department,
                "year": tt.year,
                "section": tt.section,
                "created_at": tt.created_at.strftime("%Y-%m-%d %H:%M")
                if tt.created_at
                else "",
                "is_active": tt.is_active,
                "entries_count": len(entries),
                "versions": versions_list,
            }
        )

    return response


@router.delete("/{timetable_id}")
async def delete_timetable(
    timetable_id: int,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Timetable).where(Timetable.id == timetable_id))
    timetable = result.scalar_one_or_none()
    if not timetable:
        raise HTTPException(status_code=404, detail="Timetable not found")

    await db.execute(
        delete(TimetableEntry).where(TimetableEntry.timetable_id == timetable_id)
    )
    await db.execute(
        delete(TimetableVersion).where(
            TimetableVersion.timetable_id == timetable_id
        )
    )
    await db.delete(timetable)

    await db.flush()
    return {"message": "Timetable deleted successfully"}


# ─── DCRA+ Extended Request Models ──────────────────────────────────────────

class AllocateRoomsPreviewRequest(BaseModel):
    department: str
    year: str
    section: str
    entries: List[TimetableEntrySchema]


class ManualOverrideRequest(BaseModel):
    timetable_entry_id: int
    classroom_number: str
    override_reason: str


class ReallocateEventRequest(BaseModel):
    event_type: str  # "maintenance", "locked", "emergency", "faculty_leave"
    target_id: int   # classroom_id or timetable_entry_id
    details: Optional[str] = ""


# ─── DCRA+ API Endpoints ─────────────────────────────────────────────────────

@router.post("/allocate-rooms-preview")
async def allocate_rooms_preview(
    req: AllocateRoomsPreviewRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    allocations_batch = []
    output_entries = []
    
    # Sort entries by day and start time to align consecutive walking distance optimization
    sorted_entries = sorted(req.entries, key=lambda x: (x.day, time_to_minutes(x.start_time)))
    
    for entry in sorted_entries:
        is_lunch = entry.subject.lower() in ["lunch", "lunch break", "break", "interval"]
        
        if is_lunch:
            output_entries.append({
                "day": entry.day,
                "subject": entry.subject,
                "faculty": entry.faculty or "N/A",
                "start_time": entry.start_time,
                "end_time": entry.end_time,
                "room": "N/A",
                "suitability_score": 100.0,
                "predicted_occupancy": 0,
                "occupancy_confidence": 1.0,
                "student_movement_distance": 0.0,
                "faculty_movement_distance": 0.0,
                "explanation": json.dumps({"overall_suitability": "100% - Lunch Break"}),
                "candidate_rooms": []
            })
            continue

        alloc = await DCRA_Algorithm.allocate_room_for_entry(
            db=db,
            timetable_id=0, # temp preview
            subject=entry.subject,
            faculty=entry.faculty,
            department=req.department,
            year=req.year,
            section=req.section,
            day=entry.day,
            start_time=entry.start_time,
            end_time=entry.end_time,
            current_allocations_batch=allocations_batch
        )
        
        room_name = alloc["room"]
        allocations_batch.append({
            "room": room_name,
            "day": entry.day,
            "start_time": entry.start_time,
            "end_time": entry.end_time
        })
        
        output_entries.append({
            "day": entry.day,
            "subject": entry.subject,
            "faculty": entry.faculty or "N/A",
            "start_time": entry.start_time,
            "end_time": entry.end_time,
            "room": room_name,
            "suitability_score": alloc["suitability_score"],
            "predicted_occupancy": alloc["predicted_occupancy"],
            "occupancy_confidence": alloc.get("occupancy_confidence", 1.0),
            "student_movement_distance": alloc.get("student_movement_distance", 0.0),
            "faculty_movement_distance": alloc.get("faculty_movement_distance", 0.0),
            "explanation": alloc["explanation"],
            "candidate_rooms": alloc.get("candidate_rooms", [])
        })

    return output_entries


@router.post("/{timetable_id}/allocate-rooms")
async def allocate_rooms_for_timetable(
    timetable_id: int,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    # Load Timetable
    t_res = await db.execute(select(Timetable).where(Timetable.id == timetable_id))
    timetable = t_res.scalar_one_or_none()
    if not timetable:
        raise HTTPException(status_code=404, detail="Timetable not found")

    # Load Entries
    e_res = await db.execute(select(TimetableEntry).where(TimetableEntry.timetable_id == timetable_id))
    entries = e_res.scalars().all()
    
    # Sort entries
    sorted_entries = sorted(entries, key=lambda x: (x.day, x.start_time))
    
    # Delete old ClassroomAllocations for this timetable
    entry_ids = [e.id for e in entries]
    if entry_ids:
        await db.execute(
            delete(ClassroomAllocation).where(ClassroomAllocation.timetable_entry_id.in_(entry_ids))
        )
        await db.flush()

    allocations_batch = []
    response_data = []

    for entry in sorted_entries:
        is_lunch = entry.subject.lower() in ["lunch", "lunch break", "break", "interval"]
        if is_lunch:
            entry.room = "N/A"
            response_data.append({
                "entry_id": entry.id,
                "subject": entry.subject,
                "room": "N/A",
                "suitability_score": 100.0
            })
            continue

        start_str = entry.start_time.strftime("%H:%M")
        end_str = entry.end_time.strftime("%H:%M")

        alloc = await DCRA_Algorithm.allocate_room_for_entry(
            db=db,
            timetable_id=timetable_id,
            subject=entry.subject,
            faculty=entry.faculty,
            department=timetable.department,
            year=timetable.year,
            section=timetable.section,
            day=entry.day,
            start_time=start_str,
            end_time=end_str,
            current_allocations_batch=allocations_batch
        )

        room_name = alloc["room"]
        entry.room = room_name
        
        allocations_batch.append({
            "room": room_name,
            "day": entry.day,
            "start_time": start_str,
            "end_time": end_str
        })

        # Save allocation
        if room_name != "TBD":
            c_stmt = select(Classroom).where(Classroom.room_number == room_name)
            c_res = await db.execute(c_stmt)
            classroom_obj = c_res.scalar_one_or_none()
            if classroom_obj:
                ca = ClassroomAllocation(
                    classroom_id=classroom_obj.id,
                    timetable_entry_id=entry.id,
                    suitability_score=alloc["suitability_score"],
                    explanation=alloc["explanation"],
                    predicted_occupancy=alloc["predicted_occupancy"],
                    occupancy_confidence=alloc.get("occupancy_confidence", 1.0),
                    student_movement_distance=alloc.get("student_movement_distance", 0.0),
                    faculty_movement_distance=alloc.get("faculty_movement_distance", 0.0),
                    status="active"
                )
                db.add(ca)

        response_data.append({
            "entry_id": entry.id,
            "subject": entry.subject,
            "room": room_name,
            "suitability_score": alloc["suitability_score"],
            "explanation": alloc["explanation"]
        })

    await db.flush()
    return {
        "message": f"DCRA+ completed. Allocated rooms for {len(response_data)} classes.",
        "allocations": response_data
    }


@router.post("/allocation/override")
async def manual_override_allocation(
    req: ManualOverrideRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    # 1. Fetch TimetableEntry
    e_stmt = select(TimetableEntry, Timetable).join(Timetable).where(TimetableEntry.id == req.timetable_entry_id)
    e_res = await db.execute(e_stmt)
    row = e_res.one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Timetable entry not found")
    
    entry, timetable = row
    previous_room = entry.room
    
    # 2. Fetch Classroom
    c_stmt = select(Classroom).where(Classroom.room_number == req.classroom_number)
    c_res = await db.execute(c_stmt)
    classroom_obj = c_res.scalar_one_or_none()
    if not classroom_obj:
        raise HTTPException(status_code=404, detail=f"Classroom '{req.classroom_number}' not found")
        
    # Check if target room is occupied at that slot by another active class
    start_str = entry.start_time.strftime("%H:%M")
    end_str = entry.end_time.strftime("%H:%M")
    curr_start = time_to_minutes(start_str)
    curr_end = time_to_minutes(end_str)
    
    conflict_stmt = select(TimetableEntry).join(Timetable).where(
        Timetable.is_active == True,
        TimetableEntry.day == entry.day,
        TimetableEntry.room == req.classroom_number,
        TimetableEntry.id != entry.id
    )
    conflict_res = await db.execute(conflict_stmt)
    conflicts = conflict_res.scalars().all()
    for c in conflicts:
        c_start = time_to_minutes(c.start_time.strftime("%H:%M"))
        c_end = time_to_minutes(c.end_time.strftime("%H:%M"))
        if max(curr_start, c_start) < min(curr_end, c_end):
            raise HTTPException(
                status_code=400,
                detail=f"Classroom '{req.classroom_number}' is occupied by '{c.subject}' at this time slot."
            )

    # 3. Update TimetableEntry
    entry.room = req.classroom_number
    
    # 4. Update/Create ClassroomAllocation
    ca_stmt = select(ClassroomAllocation).where(ClassroomAllocation.timetable_entry_id == entry.id)
    ca_res = await db.execute(ca_stmt)
    ca = ca_res.scalar_one_or_none()
    
    if not ca:
        ca = ClassroomAllocation(timetable_entry_id=entry.id)
        db.add(ca)
        
    ca.classroom_id = classroom_obj.id
    ca.is_manual_override = True
    ca.override_reason = req.override_reason
    ca.original_ai_room = previous_room
    ca.suitability_score = 100.0 # Force override suitability
    ca.explanation = json.dumps({
        "override_log": f"Manually overridden by administrator. Reason: {req.override_reason}",
        "original_ai_room": previous_room
    })
    
    # 5. Log to ReallocationLogs
    log = ReallocationLog(
        timetable_id=timetable.id,
        timetable_entry_id=entry.id,
        event_trigger="Manual Override",
        previous_room=previous_room,
        new_room=req.classroom_number,
        reason=req.override_reason
    )
    db.add(log)
    
    await db.flush()
    return {
        "message": f"Successfully overridden allocation. Class is moved to {req.classroom_number}.",
        "entry_id": entry.id,
        "room": req.classroom_number
    }


@router.get("/analytics/dcra")
async def get_dcra_analytics(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    # 1. Fetch all Classrooms and calculate CHI statuses
    c_res = await db.execute(select(Classroom))
    classrooms = c_res.scalars().all()
    
    chi_scores = []
    health_counts = {"Excellent": 0, "Good": 0, "Average": 0, "Needs Maintenance": 0}
    
    for r in classrooms:
        chi = ClassroomHealthIndexCalculator.calculate_chi(r)
        if chi["score"] > 0:
            chi_scores.append(chi["score"])
            health_counts[chi["status"]] = health_counts.get(chi["status"], 0) + 1
        else:
            health_counts["Needs Maintenance"] += 1
            
    avg_health = round(sum(chi_scores) / len(chi_scores), 1) if chi_scores else 100.0

    # 2. Get active bookings to calculate utilization
    stmt = select(TimetableEntry, Timetable).join(Timetable).where(Timetable.is_active == True)
    res = await db.execute(stmt)
    active_entries = res.all()
    
    TOTAL_AVAILABLE_SLOTS = 36 # 6 days * 6 slots
    utilization_data = []
    room_booking_counts = {}
    
    unused_rooms = []
    overloaded_rooms = []
    building_usage = {}
    dept_usage = {}
    total_waste_seats = 0
    
    for room in classrooms:
        room_entries = [
            (e, t) for e, t in active_entries 
            if e.room and e.room.strip().lower() == room.room_number.strip().lower()
        ]
        
        booked_slots = len(room_entries)
        room_booking_counts[room.room_number] = booked_slots
        util_rate = min(100, int((booked_slots / TOTAL_AVAILABLE_SLOTS) * 100))
        
        if booked_slots == 0:
            unused_rooms.append(room.room_number)
        if util_rate > 70:
            overloaded_rooms.append(room.room_number)
            
        building_usage[room.building] = building_usage.get(room.building, 0) + booked_slots
        
        for e, t in room_entries:
            dept_usage[t.department] = dept_usage.get(t.department, 0) + 1
            
            # Resource Waste (seat difference)
            # Find allocation
            ca_stmt = select(ClassroomAllocation).where(ClassroomAllocation.timetable_entry_id == e.id)
            ca_res = await db.execute(ca_stmt)
            ca = ca_res.scalar_one_or_none()
            predicted = ca.predicted_occupancy if ca else int(room.capacity * 0.85)
            
            waste = max(0, room.capacity - predicted)
            total_waste_seats += waste

    # Group heatmap occupancy data
    # Slots: 09:00, 10:00, 11:00, 13:00, 14:00, 15:00
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
    slots = ["09:00", "10:00", "11:00", "13:00", "14:00", "15:00"]
    heatmap = []
    
    for day in days:
        for slot in slots:
            slot_mins = time_to_minutes(slot)
            # Count rooms occupied in this slot
            occupied_count = 0
            for e, t in active_entries:
                if e.day.lower() == day.lower() and e.room and e.room != "N/A" and e.room != "TBD":
                    e_start = time_to_minutes(e.start_time.strftime("%H:%M"))
                    e_end = time_to_minutes(e.end_time.strftime("%H:%M"))
                    if e_start <= slot_mins < e_end:
                        occupied_count += 1
            
            occupancy_rate = min(100, int((occupied_count / max(1, len(classrooms))) * 100))
            heatmap.append({
                "day": day,
                "time": slot,
                "occupancy": occupancy_rate
            })

    # Fetch recent Reallocation Logs
    log_stmt = select(ReallocationLog, TimetableEntry).join(
        TimetableEntry, ReallocationLog.timetable_entry_id == TimetableEntry.id
    ).order_by(ReallocationLog.timestamp.desc()).limit(8)
    log_res = await db.execute(log_stmt)
    logs = log_res.all()
    
    recent_logs = [
        {
            "id": l.id,
            "subject": entry.subject,
            "trigger": l.event_trigger,
            "previous_room": l.previous_room or "N/A",
            "new_room": l.new_room,
            "reason": l.reason,
            "time": l.timestamp.strftime("%Y-%m-%d %H:%M")
        }
        for l, entry in logs
    ]

    return {
        "average_health": avg_health,
        "health_distribution": health_counts,
        "total_classrooms": len(classrooms),
        "unused_classrooms": unused_rooms,
        "overloaded_classrooms": overloaded_rooms,
        "building_usage": building_usage,
        "department_usage": dept_usage,
        "total_seats_wasted": total_waste_seats,
        "heatmap": heatmap,
        "recent_reallocations": recent_logs
    }


@router.post("/allocation/reallocate-event")
async def trigger_reallocate_event(
    req: ReallocateEventRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
):
    reallocations = await DynamicReallocationEngine.trigger_reallocation(
        db=db,
        event_type=req.event_type,
        target_id=req.target_id,
        details=req.details
    )
    return {
        "message": f"Successfully processed event-driven reallocation trigger. Reallocated {len(reallocations)} classes.",
        "reallocations": reallocations
    }


@router.get("/templates/download/{template_type}")
@router.get("/templates/download")
async def download_template(template_type: str = "timetable"):
    import io
    from app.services.import_service import ImportService
    
    if template_type == "timetable":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Timetable Template"
        headers = ["day", "subject", "faculty", "start_time", "end_time", "room"]
        ws.append(headers)
        ws.append(["Monday", "Data Structures", "Dr. Amit Kumar", "09:00", "11:00", "Room 101"])
        ws.append(["Monday", "Operating Systems Lab", "Dr. Sneha Gupta", "13:00", "16:00", "OS Lab"])
    else:
        wb = ImportService.create_excel_template(template_type)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    headers = {
        'Content-Disposition': f'attachment; filename="{template_type}_template.xlsx"'
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )


@router.post("/import/preview")
async def import_preview(
    file: UploadFile = File(...),
    template_type: str = Form(...),
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db)
):
    import io
    from openpyxl import load_workbook
    from app.services.import_service import ImportService
    
    contents = await file.read()
    wb = load_workbook(filename=io.BytesIO(contents), data_only=True)
    
    if template_type == "all_in_one":
        report = await ImportService.validate_and_parse_all_in_one(db, wb)
    else:
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(row)
        report = await ImportService.validate_and_parse_import(db, rows, template_type)
        
    return report


@router.post("/import/execute")
async def import_execute(
    req: dict,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db)
):
    from app.services.import_service import ImportService
    template_type = req.get("template_type")
    
    if template_type == "all_in_one":
        categories = req.get("categories", {})
        res = await ImportService.execute_import_all_in_one(db, categories, current_user["id"])
    else:
        valid_records = req.get("valid_records", [])
        res = await ImportService.execute_import(db, valid_records, template_type, current_user["id"])
        
    return res


@router.get("/import/history")
async def get_import_history(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db)
):
    history_res = await db.execute(
        select(ImportHistory).order_by(ImportHistory.created_at.desc())
    )
    records = history_res.scalars().all()
    return [
        {
            "id": r.id,
            "filename": r.filename,
            "import_type": r.import_type,
            "total_records": r.total_records,
            "successful_imports": r.successful_imports,
            "failed_imports": r.failed_imports,
            "status": r.status,
            "processing_time_ms": r.processing_time_ms,
            "created_at": r.created_at.isoformat() if r.created_at else None
        }
        for r in records
    ]


class ReassignRoomRequest(BaseModel):
    section_id: int
    room_id: int


@router.post("/classroom/reassign")
async def reassign_classroom(
    req: ReassignRoomRequest,
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db)
):
    section_stmt = select(Section).where(Section.id == req.section_id)
    section_res = await db.execute(section_stmt)
    sect = section_res.scalar_one_or_none()
    if not sect:
        raise HTTPException(status_code=404, detail="Section not found")
        
    room_stmt = select(Classroom).where(Classroom.id == req.room_id)
    room_res = await db.execute(room_stmt)
    room_obj = room_res.scalar_one_or_none()
    if not room_obj:
        raise HTTPException(status_code=404, detail="Classroom not found")
        
    sect.permanent_room_id = room_obj.id
    db.add(sect)
    await db.commit()
    return {"success": True, "message": f"Successfully reassigned permanent room to {room_obj.room_number}"}


@router.get("/sections")
async def get_sections(
    current_user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db)
):
    from sqlalchemy.orm import selectinload
    stmt = select(Section)
    res = await db.execute(stmt)
    sections = res.scalars().all()
    
    output = []
    for s in sections:
        dept_stmt = select(Department).where(Department.id == s.department_id)
        dept_res = await db.execute(dept_stmt)
        dept = dept_res.scalar_one_or_none()
        
        room_name = "None"
        if s.permanent_room_id:
            room_stmt = select(Classroom).where(Classroom.id == s.permanent_room_id)
            room_res = await db.execute(room_stmt)
            room_obj = room_res.scalar_one_or_none()
            if room_obj:
                room_name = room_obj.room_number
                
        output.append({
            "id": s.id,
            "department": dept.name if dept else "Unknown",
            "department_code": dept.code if dept else "UNK",
            "semester": s.semester,
            "section_name": s.section_name,
            "student_strength": s.student_strength,
            "permanent_room_id": s.permanent_room_id,
            "permanent_room_name": room_name
        })
    return output


@router.get("/{department}/{semester}")
async def get_timetable_by_dept_and_semester(
    department: str,
    semester: int,
    db: AsyncSession = Depends(get_db),
):
    return await TimetableController.get_timetable_by_dept_and_semester(department, semester, db)

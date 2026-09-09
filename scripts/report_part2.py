# -*- coding: utf-8 -*-
"""
Report Part 2: Cover, Bonafide, and Preliminaries
Calibrated to eliminate all blank pages and spillovers.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from report_part1 import (
    set_cell_margins, set_cell_shading, set_table_borders,
    add_p, add_bullet, add_chapter_heading, add_sec_heading,
    add_subsec_heading, add_figure, add_prelim_list, add_code_block
)
from report_metadata import (
    METADATA, ABSTRACT_PARAS, ACKNOWLEDGEMENT_PARAS,
    ABBREVIATIONS, FIGURES_LIST, TABLES_LIST, TOC_ENTRIES
)

def build_cover_page(doc):
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(8)
    p_top.paragraph_format.space_after = Pt(8)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    logo_path = os.path.join('Report', 'images', 'skct_logo.jpeg')
    if os.path.exists(logo_path):
        r_logo = p_top.add_run()
        r_logo.add_picture(logo_path, height=Inches(1.05))

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.line_spacing = 1.3
    r_title = p_title.add_run(METADATA['title'])
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(14.5)
    r_title.bold = True

    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep.paragraph_format.space_before = Pt(10)
    p_rep.paragraph_format.space_after = Pt(12)
    p_rep.paragraph_format.line_spacing = 1.15
    r_rep = p_rep.add_run("A PROJECT REPORT\n(PHASE I)")
    r_rep.font.name = 'Times New Roman'
    r_rep.font.size = Pt(13)
    r_rep.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(8)
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Submitted by")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11.5)
    r_sub.italic = True

    p_stud = doc.add_paragraph()
    p_stud.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stud.paragraph_format.space_before = Pt(2)
    p_stud.paragraph_format.space_after = Pt(16)
    p_stud.paragraph_format.line_spacing = 1.25
    for s in METADATA['students']:
        r_s = p_stud.add_run(f"{s['name']} (Register No: {s['reg_no']})\n")
        r_s.font.name = 'Times New Roman'
        r_s.font.size = Pt(12)
        r_s.bold = True

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(8)
    p_deg.paragraph_format.space_after = Pt(14)
    p_deg.paragraph_format.line_spacing = 1.15
    r_deg = p_deg.add_run("in partial fulfillment for the award of the degree\nof\nBACHELOR OF ENGINEERING\nIN\nCOMPUTER SCIENCE AND ENGINEERING")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(12)
    r_deg.bold = True

    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_before = Pt(10)
    p_col.paragraph_format.space_after = Pt(4)
    p_col.paragraph_format.line_spacing = 1.15
    r_col = p_col.add_run(f"{METADATA['college']}\n")
    r_col.font.name = 'Times New Roman'
    r_col.font.size = Pt(12.5)
    r_col.bold = True

    r_col_sub = p_col.add_run(f"{METADATA['college_sub']}\n{METADATA['location']}")
    r_col_sub.font.name = 'Times New Roman'
    r_col_sub.font.size = Pt(10.5)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(16)
    p_date.paragraph_format.space_after = Pt(0)
    r_date = p_date.add_run(METADATA['date'])
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(12)
    r_date.bold = True

def build_bonafide_certificate(doc):
    doc.add_page_break()

    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_before = Pt(16)
    p_col.paragraph_format.space_after = Pt(4)
    p_col.paragraph_format.line_spacing = 1.15
    r_col = p_col.add_run(f"{METADATA['college']}\n")
    r_col.font.name = 'Times New Roman'
    r_col.font.size = Pt(12.5)
    r_col.bold = True
    r_sub = p_col.add_run(f"{METADATA['college_sub']}\n{METADATA['location']}")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10)

    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert.paragraph_format.space_before = Pt(16)
    p_cert.paragraph_format.space_after = Pt(16)
    r_cert = p_cert.add_run("BONAFIDE CERTIFICATE")
    r_cert.font.name = 'Times New Roman'
    r_cert.font.size = Pt(13.5)
    r_cert.bold = True

    stud_str = ", ".join([f"\"{s['name']} ({s['reg_no']})\"" for s in METADATA['students']])
    cert_text = (
        f"Certified that this project report \"{METADATA['title']}\" is the bonafide work of "
        f"{stud_str} who carried out the project work Phase I under my supervision."
    )
    add_p(doc, cert_text, space_before=Pt(8), space_after=Pt(28), line_spacing=1.45)

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.rows[0].cells[0].width = Inches(3.2)
    sig_table.rows[0].cells[1].width = Inches(3.2)
    set_table_borders(sig_table, color="FFFFFF", sz="0", val="none")

    cell_left = sig_table.rows[0].cells[0]
    p_l = cell_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.line_spacing = 1.15
    p_l.add_run("SIGNATURE\n\n\n\n").bold = True
    p_l.add_run(f"{METADATA['supervisor']['name']}\n").bold = True
    p_l.add_run("SUPERVISOR\n").bold = True
    p_l.add_run(f"{METADATA['supervisor']['designation']},\n{METADATA['supervisor']['dept']},\n{METADATA['supervisor']['institution']}")

    cell_right = sig_table.rows[0].cells[1]
    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_r.paragraph_format.line_spacing = 1.15
    p_r.add_run("SIGNATURE\n\n\n\n").bold = True
    p_r.add_run(f"{METADATA['hod']['name']}\n").bold = True
    p_r.add_run("HEAD OF THE DEPARTMENT\n").bold = True
    p_r.add_run(f"{METADATA['hod']['designation']},\n{METADATA['hod']['dept']},\n{METADATA['hod']['institution']}")

    p_viva = doc.add_paragraph()
    p_viva.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_viva.paragraph_format.space_before = Pt(28)
    p_viva.paragraph_format.space_after = Pt(32)
    p_viva.paragraph_format.line_spacing = 1.25
    r_viva = p_viva.add_run(
        "Certified that the candidates were examined by us in the Project Phase I Viva-Voce examination held on "
        "____________________ at Sri Krishna College of Technology, Kovaipudur, Coimbatore - 641042."
    )
    r_viva.font.name = 'Times New Roman'
    r_viva.font.size = Pt(11)

    ex_table = doc.add_table(rows=1, cols=2)
    ex_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ex_table.rows[0].cells[0].width = Inches(3.2)
    ex_table.rows[0].cells[1].width = Inches(3.2)
    set_table_borders(ex_table, color="FFFFFF", sz="0", val="none")

    p_ex1 = ex_table.rows[0].cells[0].paragraphs[0]
    p_ex1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ex1.add_run("INTERNAL EXAMINER").bold = True

    p_ex2 = ex_table.rows[0].cells[1].paragraphs[0]
    p_ex2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ex2.add_run("EXTERNAL EXAMINER").bold = True

def setup_preliminaries_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
    sec._sectPr.append(pgNumType)

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = footer_p.add_run()
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    r_f._r.append(fld)
    return sec

def build_abstract(doc):
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(14)
    p_hdr.paragraph_format.space_after = Pt(10)
    r = p_hdr.add_run("ABSTRACT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    for p_text in ABSTRACT_PARAS:
        add_p(doc, p_text, space_before=Pt(0), space_after=Pt(3.5), line_spacing=1.18)

def build_acknowledgement(doc):
    doc.add_page_break()
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(14)
    p_hdr.paragraph_format.space_after = Pt(10)
    r = p_hdr.add_run("ACKNOWLEDGEMENT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    for p_text in ACKNOWLEDGEMENT_PARAS:
        add_p(doc, p_text, space_before=Pt(0), space_after=Pt(3), line_spacing=1.18)

def build_list_of_figures(doc):
    doc.add_page_break()
    headers = ["FIGURE NO.", "TITLE", "PAGE NO."]
    fig_rows = [[f[0], f[1], f[2]] for f in FIGURES_LIST]
    col_widths = [Inches(1.4), Inches(4.3), Inches(0.9)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT]
    add_prelim_list(doc, "LIST OF FIGURES", headers, fig_rows, col_widths, aligns)

def build_list_of_tables(doc):
    doc.add_page_break()
    headers = ["TABLE NO.", "TITLE", "PAGE NO."]
    tbl_rows = [[t[0], t[1], t[2]] for t in TABLES_LIST]
    col_widths = [Inches(1.4), Inches(4.3), Inches(0.9)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT]
    add_prelim_list(doc, "LIST OF TABLES", headers, tbl_rows, col_widths, aligns)

def build_list_of_abbreviations(doc):
    doc.add_page_break()
    headers = ["ABBREVIATION", "EXPANSION"]
    col_widths = [Inches(1.8), Inches(4.8)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
    add_prelim_list(doc, "LIST OF ABBREVIATIONS", headers, ABBREVIATIONS, col_widths, aligns)

def build_table_of_contents(doc):
    doc.add_page_break()
    headers = ["CHAPTER NO.", "TITLE", "PAGE NO."]
    col_widths = [Inches(1.4), Inches(4.3), Inches(0.9)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT]
    add_prelim_list(doc, "TABLE OF CONTENTS", headers, TOC_ENTRIES, col_widths, aligns)

print("report_part2.py updated with compact single-page prelims.")
